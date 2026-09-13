from app.models.schemas import RawPage, SectionType, Chunk, EmbeddedChunk
from typing import List
import json
import logging
from pydantic import BaseModel
from typing import Optional
from pathlib import Path
import fitz
from pptx import Presentation
import hashlib
import re
from app.clients.embedding_client import embed_batch
from app.db.vector_store import upsert_chunks
from app.config import get_settings
from app.db.vector_store import existing_hashes

logger = logging.getLogger(__name__)

MIN_PAGE_CHARS = 80
CHARS_PER_TOKEN = 4

SECTION_KEYWORDS = {
    SectionType.WARRANTY: ["warranty", "guarantee", "return policy"],
    SectionType.WARNING: ["caution", "warning", "precaution", "do not", "avoid", "side effect"],
    SectionType.USAGE: ["how to use", "direction", "apply", "usage", "routine", "step"],
    SectionType.SPEC: ["ingredient", "material code", "vol.", "ml", "spf", "pa+", "formula", "ph "],
    SectionType.CLAIM: ["clinical", "tested", "proven", "certified", "%", "free from", "dermatolog"],
}


class ManifestEntry(BaseModel):
    file: str
    product_name: str
    line: Optional[str] = None


def _load_pdf_pages(path: Path, product_name: str) -> List[RawPage]: 
    pages: List[RawPage] = []
    with fitz.open(path) as doc:
        for page_index, page in enumerate(doc):
            text = page.get_text().strip()
            if text:
                pages.append(
                    RawPage(product_name=product_name, page_number=page_index + 1, text=text)
                )
    return pages


def _load_pptx_pages(path: Path, product_name: str) -> List[RawPage]:
    pages: List[RawPage] = []
    prs = Presentation(path)
    for slide_index, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    texts.append(" | ".join(cells))
        text = "\n".join(texts).strip()

        if text:
            pages.append(
                RawPage(product_name=product_name, page_number=slide_index + 1, text = text)
            )
    return pages


def load_manuals(source_dir: str) -> List[RawPage]:
    source = Path(source_dir)
    manifest_path = source / "manifest.json"

    with open(manifest_path) as f:
        entries = [ManifestEntry(**item) for item in json.load(f)]

    all_pages: List[RawPage] = []
    for entry in entries:
        path = source / entry.file
        if not path.exists():
            logger.warning("manifest file not found, skipping: %s", path)
            continue

        suffix = path.suffix.lower()
        if suffix == ".pdf":
            pages = _load_pdf_pages(path, entry.product_name)
        elif suffix == ".pptx":
            pages = _load_pptx_pages(path, entry.product_name)
        else:
            logger.warning("unsupported format %s, skipping: %s", suffix, path.name)
            continue

        logger.info("loaded %s: %d pages", entry.product_name, len(pages))
        all_pages.extend(pages)

    return all_pages


def _classify_section(text: str) -> SectionType:
    lowered = text.lower()
    for section_type, keywords in SECTION_KEYWORDS.items():
        if any(kw in lowered for kw in keywords):
            return section_type
    return SectionType.OTHER


def _split_long_text(text: str, max_chars: int, overlap_chars: int) -> List[str]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    pieces: List[str] = []
    current = ""
    for para in paragraphs:
        candidate = f"{current}\n\n{para}" if current else para
        if len(candidate) <= max_chars:
            current = candidate
        else:
            if current:
                pieces.append(current)
            current = f"{current[-overlap_chars:]}\n\n{para}" if overlap_chars and current else para
    if current:
        pieces.append(current)
    return pieces


def chunk_document(pages: List[RawPage], chunk_size: int, chunk_overlap: int) -> List[Chunk]:
    max_chars = chunk_size * CHARS_PER_TOKEN
    overlap_chars = chunk_overlap * CHARS_PER_TOKEN
    chunks: List[Chunk] = []

    for page in pages:
        if len(page.text) < MIN_PAGE_CHARS:
            continue

        if len(page.text) <= max_chars:
            pieces = [page.text]
        else:
            pieces = _split_long_text(page.text, max_chars, overlap_chars)

        for index, piece in enumerate(pieces):
            chunks.append(
                Chunk(
                    chunk_id=f"{page.product_name}_p{page.page_number}_c{index}",
                    product_name=page.product_name,
                    section_type=_classify_section(piece),
                    page_number=page.page_number,
                    text=piece,
                    content_hash=hashlib.sha256(piece.encode("utf-8")).hexdigest(),
                )
            )

    logger.info("chunked %d pages into %d chunks", len(pages), len(chunks))
    return chunks


def embed_chunks(chunks: List[Chunk]) -> List[EmbeddedChunk]:
    if not chunks:
        return []
    vectors = embed_batch([chunk.text for chunk in chunks], task_type="RETRIEVAL_DOCUMENT")
    return [
        EmbeddedChunk(**chunk.model_dump(), vector=vector)
        for chunk, vector in zip(chunks, vectors)
    ]

def store_chunks(embedded_chunks: List[EmbeddedChunk]) -> int:
    return upsert_chunks(embedded_chunks)

def run_ingestion(source_dir: str) -> dict:
    settings = get_settings()
    pages = load_manuals(source_dir)
    chunks = chunk_document(pages, settings.chunk_size, settings.chunk_overlap)

    already = existing_hashes([c.content_hash for c in chunks])
    new_chunks = [c for c in chunks if c.content_hash not in already]
    logger.info("%d chunks total, %d already stored, %d to embed", len(chunks), len(already), len(new_chunks))

    stored = store_chunks(embed_chunks(new_chunks))
    return {
        "pages": len(pages),
        "chunks": len(chunks),
        "skipped": len(chunks) - len(new_chunks),
        "stored": stored,
    }
